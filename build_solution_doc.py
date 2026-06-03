from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUT_DIR = Path("智能客服方案文档素材")
OUT_DIR.mkdir(exist_ok=True)

DOCX_PATH = Path("智能客服与客户运营中台完整方案.docx")

FONT_PATH = "/System/Library/Fonts/STHeiti Medium.ttc"
FONT_PATH_LIGHT = "/System/Library/Fonts/STHeiti Light.ttc"
WORD_CN_FONT = "PingFang SC"
WORD_LATIN_FONT = "Calibri"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(92, 102, 112)
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
PALE_GREEN = "E8F3EF"
PALE_GOLD = "FFF4D6"
BORDER = "C7D0DA"


def font(size: int, bold: bool = False):
    try:
        return ImageFont.truetype(FONT_PATH if bold else FONT_PATH_LIGHT, size)
    except Exception:
        return ImageFont.load_default()


def set_run_font(run, size=None, bold=None, color=None, name=WORD_CN_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), WORD_LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), WORD_LATIN_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, size=None, bold=None, color=None, name=WORD_CN_FONT):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), WORD_LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), WORD_LATIN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_width(table, widths_inches):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_inches):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)
                tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(int(width * 1440)))
                tc_w.set(qn("w:type"), "dxa")


def paragraph_border_bottom(paragraph, color="B7C7D8", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)


def add_para(doc, text="", style=None, size=None, bold=None, color=None, align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_bullets(doc, items, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_run_font(r, size=10.5, color=INK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(item)
        set_run_font(r, size=10.5, color=INK)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_note_box(doc, title, body, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(table, color="D8E2EC", size="4")
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, bottom=130, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], header_fill)
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i == 0 and len(str(val)) <= 8:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val))
            set_run_font(r, size=font_size, color=INK)
    set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED)


def rounded_rect(draw, xy, radius, fill, outline=None, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_centered_text(draw, box, text, fnt, fill=(31, 41, 55), max_chars=12, line_gap=6):
    x1, y1, x2, y2 = box
    lines = []
    for part in text.split("\n"):
        wrapped = wrap(part, width=max_chars) or [part]
        lines.extend(wrapped)
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + ((x2 - x1) - w) / 2, y), line, font=fnt, fill=fill)
        y += h + line_gap


def draw_arrow(draw, start, end, color=(82, 101, 120), width=4):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        pts = [(x2, y2), (x2 - 14, y2 - 9), (x2 - 14, y2 + 9)]
    else:
        pts = [(x2, y2), (x2 + 14, y2 - 9), (x2 + 14, y2 + 9)]
    draw.polygon(pts, fill=color)


def create_lifecycle_diagram(path: Path):
    img = Image.new("RGB", (1800, 760), "white")
    d = ImageDraw.Draw(img)
    title_f = font(42, True)
    body_f = font(28, True)
    small_f = font(22)
    d.text((70, 45), "客户旅程运营闭环", font=title_f, fill=(24, 54, 83))
    d.text((70, 102), "从平台交易客户进入智能筛选、销售转化、VIP维护，再把结果回流为下一轮策略优化", font=small_f, fill=(82, 96, 110))
    boxes = [
        ("平台交易数据\n客户基础池", "#E8EEF5"),
        ("电销筛选\n外呼 + 短信", "#EAF4FF"),
        ("企微沉淀\n标签培育", "#E8F3EF"),
        ("销售承接\n会员卡转化", "#FFF4D6"),
        ("VIP私域维护\n群分流 + 报价", "#F3EAFB"),
        ("数据回流\n标签/策略优化", "#F2F4F7"),
    ]
    x, y, w, h, gap = 70, 220, 250, 150, 34
    centers = []
    for idx, (label, fill) in enumerate(boxes):
        bx = x + idx * (w + gap)
        rounded_rect(d, (bx, y, bx + w, y + h), 22, fill, outline="#B7C7D8", width=3)
        draw_centered_text(d, (bx + 18, y + 18, bx + w - 18, y + h - 18), label, body_f, max_chars=9)
        centers.append((bx + w, y + h / 2))
        if idx < len(boxes) - 1:
            draw_arrow(d, (bx + w + 8, y + h / 2), (bx + w + gap - 8, y + h / 2))
    d.arc((250, 380, 1510, 675), 0, 180, fill=(82, 101, 120), width=5)
    draw_arrow(d, (295, 527), (115, 370), color=(82, 101, 120), width=5)
    d.text((650, 605), "回复、点击、订阅、成交、分流结果全部回流客户档案", font=small_f, fill=(62, 74, 88))
    img.save(path)


def create_layer_diagram(path: Path):
    img = Image.new("RGB", (1800, 1080), "white")
    d = ImageDraw.Draw(img)
    title_f = font(42, True)
    layer_f = font(28, True)
    small_f = font(21)
    d.text((70, 45), "系统分层与协作关系", font=title_f, fill=(24, 54, 83))
    layers = [
        ("工作台与入口层", "企微侧边栏、线索池、客户360、VIP群分流台、报价页、数据看板", "#EAF4FF"),
        ("触达与执行层", "短信、企微私聊、企微群、群发、人工任务、SLA提醒", "#E8F3EF"),
        ("Agent编排层", "外呼筛选、电销培育、销售承接、VIP群分流、报价推荐、话术辅助", "#FFF4D6"),
        ("客户资产与策略层", "统一客户档案、标签规则、客户状态、模板中心、转交规则", "#F3EAFB"),
        ("数据接入与身份层", "平台交易、会员卡、报价单、外呼、短信、企微消息、统一ID映射", "#F2F4F7"),
        ("风控审计层", "权限、合规、敏感内容拦截、低置信度转人工、操作日志", "#F8E8E8"),
    ]
    x, y, w, h = 120, 145, 1560, 108
    for i, (name, desc, fill) in enumerate(layers):
        by = y + i * (h + 18)
        rounded_rect(d, (x, by, x + w, by + h), 16, fill, outline="#B7C7D8", width=3)
        d.text((x + 36, by + 26), name, font=layer_f, fill=(31, 77, 120))
        d.text((x + 390, by + 32), desc, font=small_f, fill=(56, 65, 77))
        if i < len(layers) - 1:
            draw_arrow(d, (x + w / 2, by + h + 3), (x + w / 2, by + h + 18), width=3)
    note_y = 955
    rounded_rect(d, (120, note_y - 22, 1680, note_y + 58), 12, "#FFFFFF", outline="#D8E2EC", width=2)
    d.text((150, note_y), "协作原则：底层提供统一事实，上层负责触达与执行；所有业务动作回写为事件，进入客户档案和策略优化。", font=small_f, fill=(82, 96, 110))
    img.save(path)


def create_flow_diagram(path: Path):
    img = Image.new("RGB", (1800, 1060), "white")
    d = ImageDraw.Draw(img)
    title_f = font(42, True)
    lane_f = font(28, True)
    box_f = font(23, True)
    small_f = font(19)
    d.text((70, 45), "客户、任务与团队协作流转", font=title_f, fill=(24, 54, 83))
    lanes = [
        ("电销", ["待筛选", "待外呼", "电销企微培育", "升级销售线索"], "#EAF4FF"),
        ("销售", ["销售企微承接", "需求识别", "销售跟进", "成交/继续培育"], "#FFF4D6"),
        ("私域/VIP", ["已购会员", "VIP群维护", "问题分流", "报价订阅/复购"], "#E8F3EF"),
    ]
    start_y = 155
    box_w, box_h, gap = 285, 88, 70
    for lane_idx, (lane, items, fill) in enumerate(lanes):
        y = start_y + lane_idx * 250
        d.text((75, y + 28), lane, font=lane_f, fill=(31, 77, 120))
        for i, item in enumerate(items):
            x = 260 + i * (box_w + gap)
            rounded_rect(d, (x, y, x + box_w, y + box_h), 18, fill, outline="#B7C7D8", width=3)
            draw_centered_text(d, (x + 16, y + 14, x + box_w - 16, y + box_h - 14), item, box_f, max_chars=9)
            if i < len(items) - 1:
                draw_arrow(d, (x + box_w + 8, y + box_h / 2), (x + box_w + gap - 8, y + box_h / 2), width=4)
        if lane_idx < len(lanes) - 1:
            sx = 260 + 3 * (box_w + gap) + box_w / 2
            draw_arrow(d, (sx, y + box_h + 18), (260 + box_w / 2, y + 250 - 18), width=4)
    rounded_rect(d, (220, 900, 1580, 990), 14, "#F2F4F7", outline="#B7C7D8", width=3)
    d.text((250, 928), "统一任务中心承接所有转交：销售线索、VIP问题、报价咨询、售后分流。任务处理结果再回写客户档案。", font=small_f, fill=(56, 65, 77))
    img.save(path)


def setup_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 11, False, INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        set_style_font(style, size, True, color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Bullet 2", "List Number"):
        st = styles[list_style]
        set_style_font(st, 10.5, False, INK)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run("智能客服与客户运营中台完整方案")
    set_run_font(r, size=9, color=MUTED)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    r = footer.add_run("Confidential | Page ")
    set_run_font(r, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    footer._p.append(fld_begin)
    footer._p.append(instr)
    footer._p.append(fld_end)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_cover(doc):
    add_para(doc, "方案文档", size=11, bold=True, color=MUTED, after=10)
    p = add_para(doc, "智能客服与客户运营中台完整方案", size=24, bold=True, color=RGBColor(0, 0, 0), after=4)
    p.paragraph_format.line_spacing = 1.05
    add_para(doc, "从立项背景到产品体系、业务链路、系统协作与实施优先级", size=13, color=MUTED, after=18)
    metadata = [
        ("适用团队", "电销团队、销售团队、私域团队、客服/售后、运营管理团队"),
        ("方案范围", "客户筛选、企微承接、会员卡销售、VIP群维护、报价订阅、客户档案沉淀"),
        ("版本日期", "2026年6月2日"),
        ("方案定位", "立项方案 + 产品体系方案"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        lr = p.add_run(f"{label}：")
        set_run_font(lr, size=10.5, bold=True, color=INK)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=INK)
    rule = doc.add_paragraph()
    paragraph_border_bottom(rule)
    add_note_box(
        doc,
        "核心判断",
        "本系统不是单个客服机器人，而是围绕客户生命周期建立的“客户旅程运营中台”。它用统一客户档案、事件流、Agent编排和任务中心，把电销、销售、私域、客服、售后串成一个连续闭环。",
        fill=PALE_BLUE,
    )


def build_doc():
    lifecycle_png = OUT_DIR / "01_客户旅程运营闭环.png"
    layer_png = OUT_DIR / "02_系统分层协作关系.png"
    flow_png = OUT_DIR / "03_客户任务团队协作流转.png"
    create_lifecycle_diagram(lifecycle_png)
    create_layer_diagram(layer_png)
    create_flow_diagram(flow_png)

    doc = Document()
    setup_document(doc)
    add_cover(doc)

    add_heading(doc, "1. 立项背景与目标", 1)
    add_para(
        doc,
        "平台已经沉淀了一批具备交易体量和历史采购数据的客户，但这些客户在电销、销售和私域维护阶段被不同团队分散触达，客户状态、沟通记录、意向判断和服务记录没有形成统一闭环。系统建设的核心价值，是把客户从“数据资产”转化为“可运营、可转化、可维护的客户资产”。",
    )
    add_heading(doc, "1.1 当前痛点", 2)
    add_table(
        doc,
        ["团队", "当前问题", "系统要解决的关键矛盾"],
        [
            ("电销", "外呼名单多，人工判断成本高；加企微后客户沉淀量大，难以持续精细触达。", "用智能外呼和企微培育筛选出真正值得销售承接的客户。"),
            ("销售", "销售同时承接大量企微客户，难以及时判断需求和意向；销售经验难以复制。", "让AI先完成需求识别、摘要和推荐话术，把高价值客户推到销售面前。"),
            ("私域", "VIP群内客户随意@员工，员工需要进线判断，造成错人响应和人效浪费。", "由智能助手先整理需求、标准答复和分流，减少无效进线。"),
            ("运营", "报价单以图片形式分发，客户查找困难；客户关注型号和行为无法回流。", "把报价结构化为可查询、可订阅、可追踪的客户页面。"),
        ],
        [0.85, 2.9, 2.75],
    )
    add_heading(doc, "1.2 立项目标", 2)
    add_bullets(
        doc,
        [
            "提升意向客户筛选效率：减少无效外呼和无效企微触达，把客户导向正确团队。",
            "提升会员卡销售转化：帮助销售快速理解客户画像、需求、异议点和推荐会员卡。",
            "提升VIP客户维护人效：在群内预处理问题，减少错@、重复进线和跨团队打扰。",
            "沉淀统一客户资产：把交易、企微、外呼、报价、任务和服务结果统一到客户档案。",
            "形成持续优化机制：用回复、点击、订阅、成交、分流结果反哺标签、模板和策略。",
        ],
    )

    add_heading(doc, "2. 系统到底是什么", 1)
    add_para(
        doc,
        "系统的产品定位是“客户旅程运营中台”。它不替代企微、外呼系统或交易系统，而是在这些系统之上建立统一的客户理解、智能判断、任务协作和策略回流能力。",
    )
    add_note_box(
        doc,
        "一句话定义",
        "系统从平台历史交易数据出发，识别客户当前阶段和潜在需求，在合适的企微场景中触达客户，并把需要人工处理的事项交给正确的人，最终将处理结果回流到客户档案，驱动下一轮更精准的运营。",
        fill=PALE_GREEN,
    )
    add_para(doc, "系统不是：")
    add_bullets(
        doc,
        [
            "不是只在群里自动回复的机器人。",
            "不是只记录客户资料的CRM。",
            "不是只给销售推荐话术的AI助手。",
            "不是电销、销售、私域各自独立的一组工具。",
        ],
    )
    add_para(doc, "系统是：")
    add_bullets(
        doc,
        [
            "一套围绕客户生命周期运转的统一中台。",
            "一套把交易数据、企微上下文、员工任务和报价行为连接起来的协作系统。",
            "一套让多个Agent服务不同业务阶段，但共享同一客户档案和事件流的智能运营体系。",
        ],
    )

    add_heading(doc, "3. 总体闭环链路", 1)
    doc.add_picture(str(lifecycle_png), width=Inches(6.5))
    add_caption(doc, "图1：客户从平台交易数据进入运营链路，并通过成交、服务和行为结果持续回流。")
    add_para(
        doc,
        "总体链路不是三段割裂流程，而是一条连续客户旅程。客户在电销阶段被筛选，在销售阶段被转化，在VIP阶段被维护。任何一个阶段产生的新信息，都会改变客户档案和后续触达策略。",
    )
    add_numbered(
        doc,
        [
            "平台交易数据形成客户基础池：系统根据交易体量、历史采购型号、活跃度、会员状态生成待运营客户名单。",
            "电销阶段完成初筛和企微引导：明确意向客户进入销售企微，潜在意向客户进入电销企微培育。",
            "销售阶段完成承接和转化：销售企微中AI先识别需求和异议，销售负责高价值沟通与成交。",
            "VIP阶段完成服务分流和复购运营：群助手整理需求、分配处理人，报价页沉淀客户关注型号。",
            "所有结果回流策略层：客户回复、点击、订阅、成交、未成交、问题解决结果都成为下一轮运营依据。",
        ],
    )

    add_heading(doc, "4. 系统分层与流转机制", 1)
    doc.add_picture(str(layer_png), width=Inches(6.5))
    add_caption(doc, "图2：系统由数据、客户资产、智能决策、触达执行、工作台和风控审计共同构成。")
    add_heading(doc, "4.1 六层架构", 2)
    add_table(
        doc,
        ["层级", "核心职责", "输出给上层的能力"],
        [
            ("数据接入与身份层", "接入平台交易、会员卡、报价单、外呼、短信、企微消息，并完成客户ID映射。", "统一客户身份、可追踪事件、基础事实数据。"),
            ("客户资产与策略层", "维护客户档案、标签、状态、模板、转交规则和客户生命周期。", "客户360、分层规则、可解释标签、触达策略。"),
            ("Agent编排层", "按场景调用外呼筛选、电销培育、销售承接、VIP群分流和报价推荐Agent。", "意图判断、摘要、推荐动作、任务生成。"),
            ("触达与执行层", "通过短信、企微私聊、企微群、群发和任务SLA执行动作。", "客户触达、人工任务、自动回复、SLA提醒。"),
            ("工作台与入口层", "提供企微侧边栏、线索池、客户360、VIP群分流台、报价页和数据看板。", "一线处理入口、运营配置入口、管理复盘入口。"),
            ("风控审计层", "控制权限、自动发送边界、敏感内容、低置信度转人工和操作日志。", "合规可控的自动化能力。"),
        ],
        [1.35, 3.05, 2.1],
    )
    add_heading(doc, "4.2 层级之间如何流转", 2)
    add_para(
        doc,
        "数据层提供事实，客户资产层形成客户理解，Agent层做判断和推荐，执行层完成触达或任务分配，工作台让员工处理，结果再回写到事件流。这个循环让系统越用越完整，而不是一次性配置后静态运行。",
    )
    add_bullets(
        doc,
        [
            "外部系统进入中台：交易、会员卡、企微、外呼、短信、报价数据统一进入事件流。",
            "事件流更新客户档案：每一次外呼、回复、点击、订阅、跟进都会更新客户状态和标签。",
            "Agent读取客户档案：Agent不是凭单条消息判断，而是结合客户历史、当前上下文和业务规则。",
            "任务中心承接人工动作：需要人工处理的事项以任务形式分配，而不是散落在聊天窗口里。",
            "结果回流优化策略：成交、无效、继续培育、问题解决等结果用于优化规则、模板和推荐。",
        ],
    )

    add_heading(doc, "5. 各系统之间如何协作", 1)
    add_para(
        doc,
        "该中台与现有系统的关系是“承上启下”：向下接入平台和企微等事实系统，向上提供一线工作台和管理看板。它不要求一次性替换现有系统，但建议由它承担客户档案主库和运营状态主库角色。",
    )
    add_table(
        doc,
        ["系统/渠道", "在体系中的角色", "与中台的交互"],
        [
            ("平台交易系统", "提供客户是谁、买过什么、交易体量、采购频次。", "向中台同步交易事实；接收客户标签或运营状态回写。"),
            ("会员卡系统", "提供客户是否已购卡、卡种、有效期、权益使用和续费状态。", "向中台同步会员状态；接收续费/升级线索。"),
            ("企微", "客户沟通发生地，包括私聊、电销企微、销售企微和VIP小群。", "中台读取会话与群信息，输出侧边栏建议、模板回复和转交任务。"),
            ("外呼/短信系统", "负责电销首触达和企微添加引导。", "中台提供外呼名单和话术，接收通话结果、短信发送和添加结果。"),
            ("报价系统/报价页", "把图片报价单升级成结构化查询、筛选和订阅入口。", "中台提供客户关注型号，接收访问、筛选、订阅、咨询行为。"),
            ("员工工作台", "电销、销售、私域、客服、售后的任务处理和复盘入口。", "中台统一分发任务、展示客户档案、记录处理结果。"),
        ],
        [1.3, 2.6, 2.6],
    )

    add_heading(doc, "6. 最终整合机制：四个统一", 1)
    add_para(
        doc,
        "系统最终不是靠某一个机器人或某一张页面整合，而是通过四类统一机制，把客户、沟通、任务和策略收束到同一个运营闭环里。这样电销、销售、私域、客服、售后可以分工处理，但不会分裂客户资产。",
    )
    add_table(
        doc,
        ["整合机制", "统一什么", "解决的问题"],
        [
            ("统一客户档案", "所有客户都归到同一个CustomerProfile，手机号、平台ID、企微ID、客户群ID互相映射。", "避免电销、销售、私域各自维护客户资料，导致状态不一致。"),
            ("统一事件流", "外呼、短信、企微私聊、群聊、报价访问、订阅、人工跟进都记录为InteractionEvent。", "避免只看到局部聊天记录，看不到客户完整旅程。"),
            ("统一任务中心", "销售线索、VIP问题、报价咨询、售后分流都进入HandoffTask。", "避免靠群聊@人和人工记忆协作，让需要处理的事可分配、可追踪、可复盘。"),
            ("统一策略回流", "客户回复、点击、订阅、成交、未成交、问题解决结果都回写标签、模板和Agent策略。", "避免系统只做一次性触达，无法根据真实效果持续优化。"),
        ],
        [1.25, 2.85, 2.4],
    )
    add_note_box(
        doc,
        "整合后的运转方式",
        "客户不再只属于某个环节，而是拥有统一生命周期状态；员工不再从聊天记录里找线索，而是处理系统分配的任务；策略不再依赖个人经验，而是由客户行为和成交结果持续回流更新。",
        fill=PALE_GREEN,
    )

    add_heading(doc, "7. 业务场景流程", 1)
    doc.add_picture(str(flow_png), width=Inches(6.5))
    add_caption(doc, "图3：电销、销售和私域并不是三套系统，而是同一客户旅程中的三个处理阶段。")
    add_heading(doc, "7.1 电销筛选流程", 2)
    add_numbered(
        doc,
        [
            "系统从交易数据中生成待外呼名单，优先级由交易体量、历史采购型号、活跃度和潜在会员价值决定。",
            "外呼筛选Agent结合客户画像和通话转写输出意向等级、沟通摘要和建议动作。",
            "明确意向客户收到销售企微添加引导；潜在意向客户进入电销企微培育。",
            "电销培育Agent根据采购型号和行为标签选择审核模板做定向推送。",
            "客户回复、点击或表达需求后，系统生成销售线索任务并附带摘要和推荐跟进话术。",
        ],
    )
    add_heading(doc, "7.2 销售转化流程", 2)
    add_numbered(
        doc,
        [
            "销售在企微侧边栏看到客户简档、最近交易、关注型号、会员状态和AI摘要。",
            "销售承接Agent识别客户需求、采购周期、价格敏感度、异议点和会员卡匹配度。",
            "系统根据规则+队列分配高意向客户，销售负责关键沟通和成交动作。",
            "AI提供销售话术建议，但涉及价格承诺、合同、付款、退款等内容必须由人工确认。",
            "成交、继续培育、无效、暂缓等结果回流客户档案，影响下一次推荐和触达策略。",
        ],
    )
    add_heading(doc, "7.3 VIP私域维护流程", 2)
    add_numbered(
        doc,
        [
            "已购买会员卡客户进入专属企微小群，群内包含客户、私域顾问、销售、客服和售后。",
            "VIP群分流Agent读取上下文，判断客户真实需求、问题类型、被@对象是否正确。",
            "标准问题使用审核模板自动答复；复杂问题生成任务并分流给正确角色。",
            "群内问题进入SLA管理，负责人处理完成后回填结果。",
            "每周报价通过结构化报价页推送，客户搜索、筛选、订阅和咨询行为回流档案。",
        ],
    )

    add_heading(doc, "8. 统一客户档案与状态体系", 1)
    add_para(
        doc,
        "客户档案是系统整合的核心。电销、销售、私域、客服、售后看到的是同一个客户，只是权限和任务视角不同。所有业务协作都围绕CustomerProfile和InteractionEvent展开。",
    )
    add_table(
        doc,
        ["对象", "包含内容", "用途"],
        [
            ("CustomerProfile", "平台客户ID、手机号、企微ID、群ID、负责人、交易画像、会员状态、风险标记。", "统一识别客户，承载全生命周期画像。"),
            ("InteractionEvent", "外呼、短信、企微私聊、群聊、报价访问、订阅、人工跟进等事件。", "记录客户旅程，驱动标签更新和策略回流。"),
            ("IntentSignal", "意向等级、关注产品、采购周期、异议点、置信度、是否需转人工。", "判断是否进入销售承接或继续培育。"),
            ("HandoffTask", "来源阶段、客户、转交原因、建议负责人、摘要、SLA、下一步动作。", "把需要人工处理的事项标准化分配。"),
            ("TemplateMessage", "场景、变量、审核状态、自动发送条件、禁用规则。", "控制AI自动触达边界。"),
            ("QuoteItem", "品牌、型号、配置、价格、库存/有效期、客户可见字段、订阅状态。", "支撑报价筛选、订阅和个性化推送。"),
        ],
        [1.35, 3.2, 1.95],
    )
    add_heading(doc, "8.1 客户状态机", 2)
    add_bullets(
        doc,
        [
            "客户生命周期：待筛选 → 待外呼 → 已外呼 → 电销企微培育 / 销售企微承接 → 销售跟进 → 已购会员 → VIP维护 → 续费/复购/流失预警。",
            "线索状态：新线索 → AI识别中 → 待分配 → 跟进中 → 成交 / 无效 / 继续培育 / 暂缓。",
            "VIP问题状态：群内提问 → AI识别 → 自动答复 / 人工分流 → 负责人处理中 → 已解决 / 升级处理。",
            "报价兴趣状态：未访问 → 已访问 → 已筛选 → 已订阅型号 → 已咨询 → 销售跟进。",
        ],
    )
    add_heading(doc, "8.2 首期标签策略", 2)
    add_para(doc, "首期采用规则标签优先，保证标签可解释、可运营、可复盘。AI语义标签作为补充信息，低置信度或关键标签需要人工确认。")
    add_table(
        doc,
        ["标签类型", "示例", "应用场景"],
        [
            ("品类/型号", "iPhone、华为、Mate60、iPhone13、配件。", "定向推送、报价订阅、销售推荐。"),
            ("价值分层", "高交易额、高频采购、低频大额、新客户、沉睡客户。", "外呼优先级、销售分配、VIP维护策略。"),
            ("意向分层", "明确意向、潜在意向、需培育、无效、风险客户。", "是否升级销售、是否自动触达、是否转人工。"),
            ("行为标签", "已加企微、已回复、点击报价、订阅型号、咨询会员卡。", "触发下一步任务和推荐。"),
            ("服务标签", "售前咨询、售后问题、会员权益、报价问题、投诉风险。", "VIP群分流和SLA管理。"),
        ],
        [1.25, 2.6, 2.65],
    )

    add_heading(doc, "9. Agent协作体系", 1)
    add_para(
        doc,
        "系统中的多个Agent按业务阶段分工，但不各自维护数据。所有Agent读取同一客户档案和事件流，并把输出写回任务中心或客户档案。这样既能适配不同场景，也能避免形成多个孤立机器人。",
    )
    add_table(
        doc,
        ["Agent", "输入", "输出", "动作边界"],
        [
            ("外呼筛选Agent", "交易数据、客户名单、通话转写。", "意向等级、沟通摘要、企微引导建议。", "生成建议和引导动作，不做成交承诺。"),
            ("电销培育Agent", "采购型号、企微互动、推送点击和回复。", "兴趣标签、推送建议、升级销售判断。", "仅可发送审核模板。"),
            ("销售承接Agent", "客户档案、会员卡产品、企微上下文。", "需求、异议点、推荐会员卡、话术建议。", "成交关键节点必须销售确认。"),
            ("VIP群分流Agent", "群上下文、被@对象、历史服务记录。", "问题类型、建议处理角色、优先级、回复模板。", "标准问题可自动答复，复杂问题转人工。"),
            ("报价推荐Agent", "报价数据、采购历史、订阅型号、访问行为。", "个性化报价推荐、订阅提醒、销售跟进触发。", "不直接下单或支付。"),
        ],
        [1.25, 1.75, 2.0, 1.5],
        font_size=9.2,
    )

    add_heading(doc, "10. 产品形态与页面设计", 1)
    add_heading(doc, "10.1 两个主要入口", 2)
    add_table(
        doc,
        ["入口", "使用对象", "解决什么问题", "核心能力"],
        [
            ("企微侧边栏", "销售、客服、售后、私域顾问。", "人在企微沟通现场，需要快速理解客户并选择下一步。", "客户简档、最近交易、关注型号、AI摘要、推荐回复、转交按钮。"),
            ("独立工作台", "电销、销售主管、私域主管、运营管理员、管理层。", "需要批量处理线索、配置规则、查看任务和复盘指标。", "线索池、客户360、VIP分流台、报价后台、模板中心、任务中心、数据看板。"),
        ],
        [1.25, 1.55, 2.05, 1.65],
    )
    add_heading(doc, "10.2 核心页面", 2)
    add_bullets(
        doc,
        [
            "客户360页：展示基础信息、交易概览、采购偏好、会员状态、企微关系、沟通摘要、推荐动作和风险提示。",
            "线索池：支持按意向等级、品类、交易体量、最近互动筛选，并完成分配、抢单、关闭、回访。",
            "VIP群分流台：展示待处理群问题、建议负责人、SLA、处理进度和升级状态。",
            "报价客户页：支持品牌、型号、配置、价格筛选，支持搜索、订阅、咨询入口。",
            "模板中心：管理短信、私聊、群发、群内答复、销售跟进模板和自动发送条件。",
            "数据看板：按团队、人员、阶段查看转化、人效、响应、服务质量和报价行为。",
        ],
    )

    add_heading(doc, "11. 自动化边界与风控", 1)
    add_note_box(
        doc,
        "自动发送原则",
        "首期AI自动触达只允许“审核模板 + 变量填充”。AI自由生成内容只作为员工建议，不直接发送给客户。",
        fill=PALE_GOLD,
    )
    add_table(
        doc,
        ["允许自动发送", "必须人工确认"],
        [
            ("企微添加引导、标准欢迎语、常见会员权益说明、标准报价提醒、型号订阅提醒、VIP群标准分流提示。", "价格承诺、合同/付款/退款、售后争议、投诉情绪、法律风险、低置信度识别、高价值客户关键成交节点。"),
        ],
        [3.2, 3.3],
        header_fill=LIGHT_GRAY,
    )
    add_bullets(
        doc,
        [
            "所有自动动作必须记录模板ID、变量、触发规则、置信度、发送对象和发送时间。",
            "员工只能查看自己权限范围内的客户、群和任务。",
            "涉及企微会话内容存档、客户数据使用、短信触达，需要完成合规授权和内部审计。",
            "系统必须保留低置信度转人工机制，避免AI强行回答不确定问题。",
        ],
    )

    add_heading(doc, "12. 实现优先级", 1)
    add_table(
        doc,
        ["优先级", "目标", "建设内容"],
        [
            ("P0", "打通整体闭环，让客户能跨阶段流转。", "客户档案中心、平台交易数据接入、企微客户/客户群接入、规则标签、企微侧边栏、线索池、任务中心、模板自动回复、VIP群分流、报价查询与型号订阅页。"),
            ("P1", "提升转化和运营效率。", "智能外呼接入、会员卡推荐、报价个性化推送、销售技巧知识库、团队SLA、效果看板。"),
            ("P2", "形成持续优化能力。", "话术A/B测试、客户价值预测、流失预警、成交案例自动复盘、多Agent策略优化。"),
        ],
        [0.75, 1.8, 3.95],
    )
    add_heading(doc, "12.1 P0最小可用闭环", 2)
    add_numbered(
        doc,
        [
            "从平台数据生成客户池，并为客户建立统一CustomerProfile。",
            "企微客户和客户群消息进入事件流，侧边栏能展示客户简档。",
            "系统能根据规则标签判断客户阶段，并生成销售线索或VIP分流任务。",
            "模板中心支持审核模板，AI只能自动发送模板化内容。",
            "报价页支持查询、筛选、订阅和咨询，行为回流客户档案。",
            "看板能统计从企微添加、线索升级、销售承接到VIP分流的核心指标。",
        ],
    )

    add_heading(doc, "13. 验收指标与复盘机制", 1)
    add_table(
        doc,
        ["领域", "核心指标", "复盘目的"],
        [
            ("电销", "外呼接通率、企微添加率、电销企微回复率、线索升级率。", "判断外呼名单、话术和企微引导是否有效。"),
            ("销售", "高意向承接率、首次响应时长、会员卡成交率、继续培育转化率。", "判断线索质量和销售承接效率。"),
            ("私域", "群问题识别率、正确分流率、无效员工进线下降率、SLA达成率。", "判断VIP群助手是否真正降低人效浪费。"),
            ("报价", "报价页访问率、搜索筛选使用率、型号订阅率、报价咨询转化率。", "判断结构化报价是否提升客户自助查询和销售机会。"),
            ("系统", "标签命中率、自动回复准确率、转人工准确率、审计日志完整率。", "判断系统自动化是否稳定可控。"),
        ],
        [0.9, 2.9, 2.7],
    )
    add_heading(doc, "13.1 复盘闭环", 2)
    add_bullets(
        doc,
        [
            "日维度：查看待处理任务、SLA超时、异常转人工和关键客户提醒。",
            "周维度：复盘外呼到企微、企微到销售、销售到成交、VIP分流、报价订阅数据。",
            "月维度：复盘标签规则、模板效果、销售话术、客户价值分层和团队人效变化。",
            "策略更新：将有效模板、优秀成交案例、常见问题分流规则沉淀到模板中心和知识库。",
        ],
    )

    add_heading(doc, "14. 关键假设与后续核验", 1)
    add_bullets(
        doc,
        [
            "系统采用私有化或专有云部署，客户交易数据和企微上下文在可控环境内处理。",
            "新系统作为统一客户档案主库，再向CRM/SCRM或内部系统同步必要结果。",
            "员工使用模式为“企微侧边栏 + 独立工作台”。",
            "标签首期采用规则标签优先，AI语义标签作为补充。",
            "报价页首期支持查询、筛选、订阅和咨询，不直接承接支付或下单。",
            "企微按完整权限申请，实施前核验客户联系、客户群、会话内容存档和群发能力。",
        ],
    )
    add_note_box(
        doc,
        "后续产品深化建议",
        "下一步可以继续拆成PRD级别内容：角色权限矩阵、字段字典、页面原型说明、接口清单、标签规则表、模板审核流、SLA规则和数据看板口径。",
        fill=PALE_BLUE,
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_doc()
    print(path.resolve())
